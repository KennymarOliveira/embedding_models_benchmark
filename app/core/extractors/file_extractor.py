import io
import os
import shutil
import subprocess
import tempfile
import zipfile
from xml.etree import ElementTree as ET
from zipfile import BadZipFile

import docx
import pymupdf
from docx.opc.exceptions import PackageNotFoundError


def extract_text_from_txt(content: bytes) -> str:
    """Extrai texto de arquivo TXT tentando UTF-8 e com fallback para Latin-1."""
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        try:
            return content.decode("latin-1")
        except Exception as exc:
            raise ValueError("Arquivo TXT com codificação não suportada.") from exc


def extract_text_from_pdf(content: bytes) -> str:
    """Extrai texto de arquivo PDF usando PyMuPDF."""
    try:
        with pymupdf.open(stream=content, filetype="pdf") as doc:
            text = "\n".join(page.get_text() for page in doc if page.get_text())
            return text.strip()
    except Exception as exc:
        raise ValueError(f"PDF inválido ou corrompido: {exc}") from exc


def extract_text_from_docx(content: bytes) -> str:
    """Extrai texto de arquivo DOCX (parágrafos e tabelas)."""
    try:
        doc = docx.Document(io.BytesIO(content))
    except (PackageNotFoundError, BadZipFile, KeyError) as exc:
        raise ValueError("DOCX inválido ou corrompido.") from exc

    parts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)

    return "\n".join(parts).strip()


def extract_text_from_odt(content: bytes) -> str:
    """Extrai texto de arquivo ODT (OpenDocument) lendo content.xml."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            if "content.xml" not in zf.namelist():
                raise ValueError("Arquivo ODT inválido: content.xml não encontrado.")
            xml_data = zf.read("content.xml")
            tree = ET.fromstring(xml_data)

            paragraphs = []
            for elem in tree.iter():
                tag = elem.tag.split("}")[-1]
                if tag in ("p", "h"):
                    text = "".join(elem.itertext()).strip()
                    if text:
                        paragraphs.append(text)
            return "\n".join(paragraphs).strip()
    except Exception as exc:
        raise ValueError(f"Falha ao extrair texto do documento ODT: {exc}") from exc


def extract_text_from_doc(content: bytes) -> str:
    """Extrai texto de arquivo DOC binário via LibreOffice headless ou fallback."""
    # Conversão via LibreOffice headless
    libreoffice_bin = shutil.which("libreoffice") or shutil.which("soffice")
    if libreoffice_bin:
        with tempfile.TemporaryDirectory() as tmp_dir:
            input_file = os.path.join(tmp_dir, "document.doc")
            with open(input_file, "wb") as f:
                f.write(content)

            try:
                cmd = [
                    libreoffice_bin,
                    "--headless",
                    "--convert-to",
                    "txt:Text",
                    "--outdir",
                    tmp_dir,
                    input_file,
                ]
                res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=30)
                txt_file = os.path.join(tmp_dir, "document.txt")
                if os.path.exists(txt_file):
                    with open(txt_file, "rb") as f:
                        return f.read().decode("utf-8", errors="ignore").strip()
            except Exception:
                pass

    # Fallback para extração de strings ASCII legíveis
    filtered = "".join(chr(b) for b in content if 32 <= b <= 126 or b in (10, 13))
    if len(filtered.strip()) > 20:
        return filtered.strip()

    raise ValueError("Não foi possível extrair texto do arquivo DOC binário.")


def extract_text_by_pages(filename: str, content: bytes) -> list[tuple[int, str]]:
    """Extrai o texto estruturado por página/seção [(page_num, text), ...], 1-indexed."""
    ext = filename.split(".")[-1].lower()

    if ext == "txt":
        return [(1, extract_text_from_txt(content))]
    if ext == "pdf":
        try:
            with pymupdf.open(stream=content, filetype="pdf") as doc:
                return [(i + 1, page.get_text().strip()) for i, page in enumerate(doc) if page.get_text().strip()]
        except Exception as exc:
            raise ValueError(f"PDF inválido ou corrompido: {exc}") from exc
    if ext in ["docx"]:
        return [(1, extract_text_from_docx(content))]
    if ext == "odt":
        return [(1, extract_text_from_odt(content))]
    if ext == "doc":
        return [(1, extract_text_from_doc(content))]

    raise ValueError(f"Extensão não suportada: {ext}")


def extract_text(filename: str, content: bytes) -> str:
    """Extrai todo o texto do documento a partir do nome e bytes do arquivo."""
    ext = filename.split(".")[-1].lower()

    if ext == "txt":
        return extract_text_from_txt(content)
    if ext == "pdf":
        return extract_text_from_pdf(content)
    if ext == "docx":
        return extract_text_from_docx(content)
    if ext == "odt":
        return extract_text_from_odt(content)
    if ext == "doc":
        return extract_text_from_doc(content)

    raise ValueError(f"Extensão não suportada: .{ext}. Extensões válidas: pdf, docx, doc, odt, txt")

