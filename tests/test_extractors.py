import io
import zipfile
import pytest
import pymupdf
import docx

from app.core.extractors.file_extractor import (
    extract_text,
    extract_text_by_pages,
    extract_text_from_txt,
    extract_text_from_pdf,
    extract_text_from_docx,
    extract_text_from_odt,
)


def test_extract_text_from_txt():
    content = "Este é um texto de teste em UTF-8.".encode("utf-8")
    result = extract_text("documento.txt", content)
    assert "Este é um texto de teste em UTF-8." in result


def test_extract_text_from_txt_latin1():
    content = "Texto com acentuação específica: ação e jurisdição.".encode("latin-1")
    result = extract_text("documento.txt", content)
    assert "Texto com acentuação específica" in result


def test_extract_text_from_pdf():
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Conteúdo da primeira página do PDF de teste.")
    pdf_bytes = doc.tobytes()
    doc.close()

    result = extract_text("documento.pdf", pdf_bytes)
    assert "Conteúdo da primeira página" in result

    pages = extract_text_by_pages("documento.pdf", pdf_bytes)
    assert len(pages) == 1
    assert pages[0][0] == 1
    assert "Conteúdo da primeira página" in pages[0][1]


def test_extract_text_from_docx():
    doc = docx.Document()
    doc.add_paragraph("Primeiro parágrafo do documento Word.")
    doc.add_paragraph("Segundo parágrafo com mais detalhes.")

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    result = extract_text("documento.docx", docx_bytes)
    assert "Primeiro parágrafo" in result
    assert "Segundo parágrafo" in result


def test_extract_text_from_odt():
    odt_xml = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
                         xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">
  <office:body>
    <office:text>
      <text:h>Título ODT</text:h>
      <text:p>Parágrafo extraído com sucesso do ODT.</text:p>
    </office:text>
  </office:body>
</office:document-content>"""

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("content.xml", odt_xml)
    odt_bytes = buf.getvalue()

    result = extract_text("documento.odt", odt_bytes)
    assert "Título ODT" in result
    assert "Parágrafo extraído com sucesso" in result


def test_extract_text_unsupported_extension():
    with pytest.raises(ValueError, match="Extensão não suportada"):
        extract_text("arquivo.xyz", b"qualquer conteudo")

